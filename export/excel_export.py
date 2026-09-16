# export\excel_export.py

import os
import pandas as pd
from utils.path_utils import get_desktop_dir
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.styles import Alignment

def export_sheets_to_excel(
    sheets_dict,
    filename=None,
    filepath=get_desktop_dir(),
):
    if filename is None:
        filename = os.path.join(filepath, "report.xlsx")

    os.makedirs(os.path.dirname(filename), exist_ok=True)

    with pd.ExcelWriter(
        filename,
        engine="openpyxl",
        mode="w",
    ) as writer:
        for sheet_name, df in sheets_dict.items():
            df.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )

            worksheet = writer.sheets[sheet_name]

            if sheet_name.startswith("kopas "):
                format_kopas_sheet(worksheet)

    print(f"Exported {list(sheets_dict.keys())} to {filename}")

    return filename

def format_kopas_sheet(worksheet):
    steel_blue = "4472C4"
    steel_blue_80 = "D9E1F2"
    steel_blue_40 = "8FAADC"

    header_fill = PatternFill(
        fill_type="solid",
        fgColor=steel_blue,
    )

    alternate_fill = PatternFill(
        fill_type="solid",
        fgColor=steel_blue_80,
    )

    horizontal_side = Side(
        style="thin",
        color=steel_blue_40,
    )

    outer_side = Side(
        style="thin",
        color=steel_blue_40,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

    for row in range(2, worksheet.max_row + 1):
        worksheet.cell(row=row, column=1).number_format = "DD-MMM-YYYY"
        if row % 2 == 0:
            for column in range(1, 6):
                worksheet.cell(row=row, column=column).fill = alternate_fill

    for row in range(2, worksheet.max_row + 1):
        url = worksheet.cell(row=row, column=5).value
        link_cell = worksheet.cell(row=row, column=4)

        if url:
            link_cell.hyperlink = url
            link_cell.value = "OPEN"
            link_cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            link_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=steel_blue,
            )
            link_cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    for row in range(1, worksheet.max_row + 1):
        for column in range(1, 6):
            cell = worksheet.cell(row=row, column=column)

            cell.border = Border(
                bottom=horizontal_side,
            )

    max_row = worksheet.max_row
    max_column = worksheet.max_column

    for row in range(1, max_row + 1):
        worksheet.cell(row=row, column=1).border = Border(
            left=outer_side,
            bottom=horizontal_side,
        )

        worksheet.cell(row=row, column=max_column).border = Border(
            right=outer_side,
            bottom=horizontal_side,
        )

    for column in range(1, max_column + 1):
        worksheet.cell(row=1, column=column).border = Border(
            top=outer_side,
            left=(
                outer_side
                if column == 1
                else Side(style=None)
            ),
            right=(
                outer_side
                if column == max_column
                else Side(style=None)
            ),
            bottom=horizontal_side,
        )

        worksheet.cell(row=max_row, column=column).border = Border(
            bottom=outer_side,
            left=(
                outer_side
                if column == 1
                else Side(style=None)
            ),
            right=(
                outer_side
                if column == max_column
                else Side(style=None)
            ),
        )

    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 8
    worksheet.column_dimensions["C"].width = 71
    worksheet.column_dimensions["D"].width = 10
    worksheet.column_dimensions["E"].width = 130

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.zoomScale = 115